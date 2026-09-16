"""
Geração do Excel do extrato de DACTEs (Documento Auxiliar do CT-e).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1B2A4A")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)

COLUNAS = [
    ("arquivo", "Arquivo", 26, None),
    ("numero_cte", "Número CT-e", 14, None),
    ("data_emissao", "Data de Emissão", 18, None),
    ("chave_acesso", "Chave de Acesso", 32, "@"),
    ("valor_a_receber", "Valor a Receber", 16, "#,##0.00"),
    ("aliquota_icms", "Alíquota ICMS", 14, "0.00%"),
]

COL_VALOR_A_RECEBER = 5
COL_ALIQUOTA_ICMS = 6


def gerar_excel(linhas: list[dict], avisos: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DACTE"

    for col, (_, titulo, largura, _) in enumerate(COLUNAS, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.freeze_panes = "A2"

    row = 2
    total_valor_a_receber = 0.0
    for linha in linhas:
        for col, (chave, _, _, formato) in enumerate(COLUNAS, start=1):
            valor = linha.get(chave)
            if chave == "aliquota_icms" and valor is not None:
                valor = valor / 100.0
            cell = ws.cell(row=row, column=col, value=valor)
            if formato:
                cell.number_format = formato
        total_valor_a_receber += linha.get("valor_a_receber") or 0.0
        row += 1

    total_row = row
    ws.cell(row=total_row, column=1, value=f"TOTAL — {len(linhas)} DACTE(s)").font = Font(bold=True)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=COL_VALOR_A_RECEBER - 1)
    tot_cell = ws.cell(row=total_row, column=COL_VALOR_A_RECEBER, value=round(total_valor_a_receber, 2))
    tot_cell.font = Font(bold=True)
    tot_cell.number_format = "#,##0.00"

    if avisos:
        ws2 = wb.create_sheet("Avisos")
        ws2.cell(row=1, column=1, value="Avisos do processamento").font = Font(bold=True, size=12)
        for i, aviso in enumerate(avisos, start=3):
            ws2.cell(row=i, column=1, value=aviso)
        ws2.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
