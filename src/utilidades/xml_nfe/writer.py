"""
Geração do Excel do extrato de XMLs de NF-e (uma linha por nota).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1B2A4A")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)

# (chave, título, largura, formato) — CNPJ como texto para não perder zeros à esquerda.
COLUNAS = [
    ("arquivo", "Arquivo", 34, None),
    ("numero_nf", "Número NF", 12, None),
    ("data_emissao", "Data de Emissão", 16, None),
    ("cnpj_emitente", "CNPJ Emitente", 18, "@"),
    ("cnpj_destinatario", "CNPJ Destinatário", 18, "@"),
    ("cfop", "CFOP", 12, "@"),
    ("valor_total_nota", "Valor Total da Nota", 18, "#,##0.00"),
    ("base_calculo_icms", "Base de Cálculo ICMS", 20, "#,##0.00"),
    ("valor_icms", "Valor ICMS", 14, "#,##0.00"),
]
CHAVES_SOMA = ("valor_total_nota", "base_calculo_icms", "valor_icms")


def gerar_excel(linhas: list[dict], avisos: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "XML NFE"

    for col, (_, titulo, largura, _) in enumerate(COLUNAS, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.freeze_panes = "A2"

    row = 2
    for linha in linhas:
        for col, (chave, _, _, formato) in enumerate(COLUNAS, start=1):
            cell = ws.cell(row=row, column=col, value=linha.get(chave))
            if formato:
                cell.number_format = formato
        row += 1

    primeira_soma = next(i for i, c in enumerate(COLUNAS, start=1) if c[0] in CHAVES_SOMA)
    ws.cell(row=row, column=1, value=f"TOTAL — {len(linhas)} NF-e").font = Font(bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=primeira_soma - 1)
    for col, (chave, _, _, _) in enumerate(COLUNAS, start=1):
        if chave in CHAVES_SOMA:
            cell = ws.cell(row=row, column=col, value=round(sum(l.get(chave) or 0.0 for l in linhas), 2))
            cell.font = Font(bold=True)
            cell.number_format = "#,##0.00"

    if avisos:
        ws2 = wb.create_sheet("Avisos")
        ws2.cell(row=1, column=1, value="Avisos do processamento").font = Font(bold=True, size=12)
        for i, aviso in enumerate(avisos, start=3):
            ws2.cell(row=i, column=1, value=aviso)
        ws2.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
